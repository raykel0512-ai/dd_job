from io import BytesIO
from docx import Document


def make_docx(title: str, sections: dict[str, str]) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for section_title, content in sections.items():
        doc.add_heading(section_title, level=2)
        for line in content.splitlines():
            if line.strip():
                doc.add_paragraph(line)
            else:
                doc.add_paragraph("")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
